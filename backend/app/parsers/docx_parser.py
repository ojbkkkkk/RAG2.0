from __future__ import annotations

import asyncio
from pathlib import Path

from .base import BaseParser, ParsedDocument


class DocxParser(BaseParser):
    """Word/DOCX 文档解析器（使用 python-docx）"""

    # 标题样式名称
    HEADING_STYLES = {"Heading 1", "Heading 2", "标题 1", "标题 2"}

    def supported_extensions(self) -> list[str]:
        return [".docx", ".doc"]

    async def parse(self, file_path: Path) -> ParsedDocument:
        """解析 DOCX 文件"""
        suffix = file_path.suffix.lower()

        # .doc 旧格式不支持
        if suffix == ".doc":
            return ParsedDocument(
                content="",
                metadata={
                    "error": "不支持旧版 .doc 格式，请将文件转换为 .docx 格式后重试",
                    "file_name": file_path.name,
                },
                sections=[],
                file_type="docx",
            )

        try:
            result = await asyncio.to_thread(self._parse_sync, file_path)
            return result
        except ImportError:
            return ParsedDocument(
                content="",
                metadata={"error": "python-docx 未安装，请运行: pip install python-docx"},
                sections=[],
                file_type="docx",
            )
        except Exception as e:
            return ParsedDocument(
                content="",
                metadata={"error": f"解析 DOCX 失败: {e}"},
                sections=[],
                file_type="docx",
            )

    def _parse_sync(self, file_path: Path) -> ParsedDocument:
        """同步解析 DOCX"""
        from docx import Document

        doc = Document(str(file_path))

        metadata = self._extract_metadata(doc, file_path)

        all_text_parts: list[str] = []
        sections: list[dict] = []
        current_title = ""
        current_lines: list[str] = []

        for element in doc.element.body:
            # 处理段落
            if element.tag.endswith("}p"):
                paragraph = self._find_paragraph_by_element(doc.paragraphs, element)
                if paragraph is None:
                    continue

                style_name = paragraph.style.name if paragraph.style else ""
                text = paragraph.text.strip()

                if style_name in self.HEADING_STYLES and text:
                    # 保存前一个 section
                    if current_lines or current_title:
                        section_content = self.clean_text("\n".join(current_lines))
                        if section_content or current_title:
                            sections.append({
                                "title": current_title,
                                "content": section_content,
                            })
                    current_title = text
                    current_lines = []
                    all_text_parts.append(f"\n{'#' if '1' in style_name else '##'} {text}")
                else:
                    if text:
                        current_lines.append(text)
                        all_text_parts.append(text)

            # 处理表格
            elif element.tag.endswith("}tbl"):
                table = self._find_table_by_element(doc.tables, element)
                if table is not None:
                    table_text = self._format_table(table)
                    if table_text:
                        current_lines.append(table_text)
                        all_text_parts.append(table_text)

        # 保存最后一个 section
        if current_lines or current_title:
            section_content = self.clean_text("\n".join(current_lines))
            if section_content or current_title:
                sections.append({
                    "title": current_title,
                    "content": section_content,
                })

        # 如果没有任何标题分节，将整个内容作为一个 section
        if not sections and all_text_parts:
            content = self.clean_text("\n".join(all_text_parts))
            sections.append({
                "title": "",
                "content": content,
            })

        content = self.clean_text("\n".join(all_text_parts))

        metadata["paragraph_count"] = len(doc.paragraphs)

        return ParsedDocument(
            content=content,
            metadata=metadata,
            sections=sections,
            file_type="docx",
        )

    @staticmethod
    def _find_paragraph_by_element(paragraphs, element):
        """根据 XML 元素找到对应的 Paragraph 对象"""
        for para in paragraphs:
            if para._element is element:
                return para
        return None

    @staticmethod
    def _find_table_by_element(tables, element):
        """根据 XML 元素找到对应的 Table 对象"""
        for table in tables:
            if table._element is element:
                return table
        return None

    @staticmethod
    def _format_table(table) -> str:
        """将表格转为可读文本格式"""
        lines: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append(" | ".join(cells))
        if lines:
            return "\n".join(lines)
        return ""

    @staticmethod
    def _extract_metadata(doc, file_path: Path) -> dict:
        """提取 DOCX 元数据"""
        metadata: dict = {
            "file_name": file_path.name,
        }

        core_props = doc.core_properties
        if core_props.title:
            metadata["title"] = core_props.title
        if core_props.author:
            metadata["author"] = core_props.author
        if core_props.created:
            metadata["created"] = str(core_props.created)
        if core_props.modified:
            metadata["modified"] = str(core_props.modified)
        if core_props.subject:
            metadata["subject"] = core_props.subject

        return metadata
