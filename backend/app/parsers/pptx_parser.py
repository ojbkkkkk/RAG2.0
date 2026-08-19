from __future__ import annotations

import asyncio
from pathlib import Path

from .base import BaseParser, ParsedDocument


class PptxParser(BaseParser):
    """PPT/PPTX 文档解析器（使用 python-pptx）"""

    def supported_extensions(self) -> list[str]:
        return [".pptx", ".ppt"]

    async def parse(self, file_path: Path) -> ParsedDocument:
        """解析 PPTX 文件"""
        suffix = file_path.suffix.lower()

        # .ppt 旧格式不支持
        if suffix == ".ppt":
            return ParsedDocument(
                content="",
                metadata={
                    "error": "不支持旧版 .ppt 格式，请将文件转换为 .pptx 格式后重试",
                    "file_name": file_path.name,
                },
                sections=[],
                file_type="pptx",
            )

        try:
            result = await asyncio.to_thread(self._parse_sync, file_path)
            return result
        except ImportError:
            return ParsedDocument(
                content="",
                metadata={"error": "python-pptx 未安装，请运行: pip install python-pptx"},
                sections=[],
                file_type="pptx",
            )
        except Exception as e:
            return ParsedDocument(
                content="",
                metadata={"error": f"解析 PPTX 失败: {e}"},
                sections=[],
                file_type="pptx",
            )

    def _parse_sync(self, file_path: Path) -> ParsedDocument:
        """同步解析 PPTX"""
        from pptx import Presentation

        prs = Presentation(str(file_path))

        slide_count = len(prs.slides)
        metadata: dict = {
            "slide_count": slide_count,
            "file_name": file_path.name,
        }

        all_text_parts: list[str] = []
        sections: list[dict] = []

        for slide_idx, slide in enumerate(prs.slides):
            slide_number = slide_idx + 1
            slide_title = self._extract_slide_title(slide)
            slide_text_parts: list[str] = []
            notes_text = ""

            # 提取所有形状中的文本和表格
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_text_parts.append(text)

                if shape.has_table:
                    table_text = self._format_table(shape.table)
                    if table_text:
                        slide_text_parts.append(table_text)

            # 提取备注
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                if notes_slide.notes_text_frame:
                    notes_text = notes_slide.notes_text_frame.text.strip()

            # 格式化幻灯片内容
            section_content_parts: list[str] = []

            header = f"## Slide {slide_number}: {slide_title}" if slide_title else f"## Slide {slide_number}"
            section_content_parts.append(header)

            if slide_text_parts:
                section_content_parts.append("\n".join(slide_text_parts))

            if notes_text:
                section_content_parts.append(f"备注: {notes_text}")

            section_content = self.clean_text("\n".join(section_content_parts))

            sections.append({
                "title": slide_title or f"Slide {slide_number}",
                "content": section_content,
                "slide_number": slide_number,
            })

            all_text_parts.append(section_content)

        content = self.clean_text("\n\n".join(all_text_parts))

        # 提取文档标题
        if sections and sections[0]["title"] != "Slide 1":
            metadata["title"] = sections[0]["title"]

        return ParsedDocument(
            content=content,
            metadata=metadata,
            sections=sections,
            file_type="pptx",
        )

    @staticmethod
    def _extract_slide_title(slide) -> str:
        """提取幻灯片标题"""
        if slide.shapes.title:
            return slide.shapes.title.text.strip()
        # 尝试从其他形状中寻找标题
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()
                if text and len(text) < 100:
                    return text
        return ""

    @staticmethod
    def _format_table(table) -> str:
        """将表格转为可读文本格式"""
        lines: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append(" | ".join(cells))
        if lines:
            return "\n".join(lines)
        return ""
